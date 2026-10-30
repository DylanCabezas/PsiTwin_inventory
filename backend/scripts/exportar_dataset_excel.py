"""Low-memory XLSX fallback: artifact-tool exhausted its heap on this dataset."""
import json
import sys
from datetime import datetime
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parent))
from generar_dataset_pruebas import OUT
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

def main():
    data=json.loads((OUT/'dataset.json').read_text(encoding='utf-8'))
    wb=Workbook();wb.remove(wb.active)
    for name,rows in data.items():
        sh=wb.create_sheet(name);keys=list(rows[0]);sh.append(keys)
        for r in rows:
            sh.append([datetime.fromisoformat(r[k]) if k in ('fecha','corte_entrenamiento') else r[k] for k in keys])
        for cell in sh[1]:cell.font=Font(name='Calibri',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163D50');cell.alignment=Alignment(wrap_text=True)
        sh.row_dimensions[1].height=32;sh.freeze_panes='A2'
        for i,k in enumerate(keys,1):
            sh.column_dimensions[get_column_letter(i)].width=min(65,max(20,len(k)+3))
            if k in ('fecha','corte_entrenamiento') or k.startswith('precio'):
                for cells in sh.iter_rows(min_row=2,min_col=i,max_col=i):cells[0].number_format='dd/mm/yyyy' if k in ('fecha','corte_entrenamiento') else '"S/ "0.00'
        table=Table(displayName='T'+name.replace('_',''),ref=sh.dimensions);table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);sh.add_table(table)
        if name=='README':
            sh.row_dimensions[2].height=150
            for cell in sh[2]:cell.alignment=Alignment(wrap_text=True,vertical='top')
    path=OUT/'PsiTwin_Dataset_Pruebas_2025_2026.xlsx';wb.save(path)
    check=load_workbook(path,read_only=True,data_only=True)
    assert len(check.sheetnames)==14
    assert isinstance(check['Demanda_Historica']['B2'].value,datetime)
    check.close();print(path)
if __name__=='__main__':main()
