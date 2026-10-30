"""Contract tests against generated fixture, saved XLSX and real evaluation artifacts."""
import json
import sys
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
import pytest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from generar_dataset_pruebas import build, OUT, ORIGIN, TENANT, START, END
from importar_dataset_pruebas import validate, read_excel, import_dataset

@pytest.fixture(scope='module')
def data():return build()

@pytest.mark.parametrize('sheet,count',[('Productos',30),('Empresa',1),('Almacenes',1),('Zonas',3),('Racks',6),('Niveles',10),('Ubicaciones_Productos',30),('Demanda_Historica',11880),('Pronosticos_Referencia',840)])
def test_counts(data,sheet,count):assert len(data[sheet])==count

def test_integrity(data):validate(data)
def test_reproducible(data):assert data==build()
def test_four_ice_creams(data):
    ice=[p for p in data['Productos'] if p['familia']=='Helados']
    assert len(ice)==4 and all(p['unidad']=='unidad' and p['categoria']=='Congelados' for p in ice)
def test_two_racks_per_zone(data):assert Counter(r['id_zona'] for r in data['Racks'])=={'A':2,'B':2,'C':2}
def test_no_sensors(data):assert not any('sensor' in k.casefold() for k in data)
def test_origin(data):assert all(r['id_empresa']==TENANT and r['origen']==ORIGIN and r['sintetico'] is True for rows in data.values() for r in rows)
def test_dates_continuous(data):
    expected={str(START+timedelta(days=i)) for i in range(396)}
    for p in data['Productos']:assert {r['fecha'] for r in data['Demanda_Historica'] if r['id_producto']==p['id_producto']}==expected
def test_patterns(data):
    curves=[]
    for p in data['Productos']:
        rows=[r for r in data['Demanda_Historica'] if r['id_producto']==p['id_producto']]
        qty=[r['cantidad'] for r in rows];assert 0 in qty and len(set(qty))>5;curves.append(tuple(qty))
        if p['familia']=='Helados':
            warm=[r['cantidad'] for r in rows if int(r['fecha'][5:7]) in (12,1,2,3)]
            cold=[r['cantidad'] for r in rows if int(r['fecha'][5:7]) in (6,7,8)]
            assert sum(warm)/len(warm)>sum(cold)/len(cold)*1.5
    assert len(set(curves))==30
def test_sales_linked(data):
    sales={(r['id_producto'],r['fecha']):r['cantidad'] for r in data['Movimientos'] if r['tipo_movimiento']=='SALIDA'}
    for d in data['Demanda_Historica']:assert abs(sales.get((d['id_producto'],d['fecha']),0)-d['cantidad_atendida'])<1e-8
def test_unique_movement_identifiers(data):assert len({r['id_movimiento'] for r in data['Movimientos']})==len(data['Movimientos'])
def test_no_negative_stock(data):assert all(r['stock_resultante']>=0 for r in data['Movimientos'])
def test_stock_scenarios(data):
    assert data['Inventario_Actual'][0]['stock_final']==0
    assert data['Inventario_Actual'][2]['stock_final']>data['Productos'][2]['stock_maximo']
def test_excel_actual_dates():
    from openpyxl import load_workbook
    wb=load_workbook(OUT/'PsiTwin_Dataset_Pruebas_2025_2026.xlsx',read_only=True)
    assert len(wb.sheetnames)==14;assert isinstance(wb['Demanda_Historica']['B2'].value,datetime);wb.close()
def test_excel_roundtrip(data):assert read_excel(OUT/'PsiTwin_Dataset_Pruebas_2025_2026.xlsx')==data
def test_preview_no_database(data):
    import asyncio
    assert asyncio.run(import_dataset(data))['modo']=='PREVIEW_SIN_ESCRITURAS'
def test_reject_wrong_tenant(data):
    import copy
    bad=copy.deepcopy(data);bad['Productos'][0]['id_empresa']='real'
    with pytest.raises(ValueError):validate(bad)
def test_temporal_validation_metrics():
    import csv,math
    report=json.loads((OUT/'evaluacion_autogluon.json').read_text())
    assert report['corte']<report['validacion_inicio'] and report['filas_validacion']==840
    rows=list(csv.DictReader((OUT/'predicciones_autogluon.csv').open()))
    for pid,scores in report['metricas_por_producto'].items():
        selected=[r for r in rows if r['item_id']==pid];assert len(selected)==28
        assert min(r['timestamp'] for r in selected)=='2026-10-01'
        e=[abs(float(r['real'])-float(r['mean'])) for r in selected]
        assert scores['MAE']==pytest.approx(sum(e)/28)
        assert scores['RMSE']==pytest.approx(math.sqrt(sum(v*v for v in e)/28))
        assert scores['WAPE']==pytest.approx(sum(e)/sum(float(r['real']) for r in selected)*100)
        smape=sum(200*abs(float(r['real'])-float(r['mean']))/(abs(float(r['real']))+abs(float(r['mean']))) if abs(float(r['real']))+abs(float(r['mean'])) else 0 for r in selected)/28
        assert scores['sMAPE']==pytest.approx(smape)
def test_replenishment():
    rows=json.loads((OUT/'reposicion.json').read_text(encoding='utf-8'))
    assert len(rows)==30
    for r in rows:
        assert 0<=r['reposicion_sugerida']<=max(0,r['maximo']-r['stock'])+.01
        assert r['inicio_pronostico']>'2026-10-28'
        if r['unidad']=='unidad':assert r['reposicion_sugerida']==int(r['reposicion_sugerida'])
def test_no_reference_leakage():
    source=(Path(__file__).resolve().parents[1]/'scripts/evaluar_dataset_pruebas.py').read_text()
    assert "data['Pronosticos_Referencia']" not in source
